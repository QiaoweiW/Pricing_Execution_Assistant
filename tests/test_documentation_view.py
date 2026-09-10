"""The manual must describe the app that actually exists.

Documentation rots silently: nothing fails when a page is renamed, a section
is retired or the sidebar is reordered, so the guide quietly starts lying.
This file pins the parts that can be checked mechanically against the source
of truth rather than against a copy of the prose:

* page names and their ORDER come from ``streamlit_app`` itself;
* retired features must not still be described;
* every section the guides promise must exist in the view that owns it;
* the Word export walks the same tuples the page renders, so it cannot drift.

The drift this caught when written: the guides listed pages in a different
order from the sidebar while claiming "in the same order", still told people
that Demand Summary (APS) is where you "review by corporate group and push a
patch" (retired), and never mentioned Demand Summary (IBP) at all.
"""
from __future__ import annotations

import ast
import re
from pathlib import Path

import pytest

from tests.streamlit_stub import install

install()                       # before importing anything that pulls streamlit

import pages.documentation_view as doc          # noqa: E402
from utils import docx_export                   # noqa: E402


_ROOT = Path(__file__).resolve().parent.parent
_DPA = _ROOT / "pages" / "demand_planner_analytics_view.py"


def _docx_bytes() -> bytes:
    """Build the manual, bypassing ``st.cache_data`` when it is real.

    The headless stub makes ``cache_data`` a pass-through (so there is no
    ``__wrapped__``); under real Streamlit there is one.  Handle both rather
    than assuming which module imported the page first.
    """
    fn = doc._build_manual_docx
    return getattr(fn, "__wrapped__", fn)()


def _sidebar_order() -> list[str]:
    """The real sidebar order, computed the way streamlit_app computes it."""
    src = (_ROOT / "streamlit_app.py").read_text(encoding="utf-8")
    mapping = dict(re.findall(r'"(\w+_view)":\s*"([^"]+)"', src))
    present = {p.stem for p in (_ROOT / "pages").glob("*_view.py")}
    available = [label for stem, label in mapping.items() if stem in present]
    landing = re.search(r'LANDING_VIEW: str = "([^"]+)"', src).group(1)
    pinned = re.findall(
        r'"([^"]+)"',
        re.search(r"NAV_PINNED_LAST: tuple = \(([^)]*)\)", src).group(1))
    middle = sorted(v for v in available if v != landing and v not in pinned)
    return [landing] + middle + [p for p in pinned if p in available]


# ── The guides match the sidebar ────────────────────────────────────────────

def test_every_sidebar_page_has_a_guide_in_sidebar_order():
    """The page claims "in the same order" — hold it to that."""
    assert [g.name for g in doc._PAGE_GUIDES] == _sidebar_order()


def test_no_guide_describes_a_page_that_does_not_exist():
    documented = {g.name for g in doc._PAGE_GUIDES}
    assert documented <= set(_sidebar_order())


def test_every_guide_is_actually_filled_in():
    for g in doc._PAGE_GUIDES:
        assert g.one_liner.strip() and len(g.steps) >= 3, g.name
        assert all(str(s).strip() for s in g.steps), g.name


# ── Retired features must not still be documented ───────────────────────────

@pytest.mark.parametrize("phrase", [
    "corporate group, and push a patch",
    "review it by corporate group",
    "Update Price Books",
    "Weekly & Monthly Butter Movers",
])
def test_retired_features_are_not_described(phrase):
    text = (_ROOT / "pages" / "documentation_view.py").read_text(encoding="utf-8")
    assert phrase not in text, f"the manual still describes retired: {phrase}"


def test_the_aps_guide_describes_two_uploads_not_a_patch_workflow():
    aps = next(g for g in doc._PAGE_GUIDES if g.name == "Demand Planner Analytics")
    body = " ".join(aps.steps)
    assert "RO_Seed" in body and "APS bulk export" in body
    assert "patch" not in body.lower()


# ── The guides describe sections that exist ─────────────────────────────────

def _expander_titles(path: Path) -> list[str]:
    tree = ast.parse(path.read_text(encoding="utf-8"))
    out = []
    for node in ast.walk(tree):
        if (isinstance(node, ast.Call) and isinstance(node.func, ast.Attribute)
                and node.func.attr == "expander" and node.args
                and isinstance(node.args[0], ast.Constant)
                and isinstance(node.args[0].value, str)):
            out.append(node.args[0].value)
    return out


@pytest.mark.parametrize("section", [
    "IBP Cadence and Supporting files",
    "Business Health",
    "RO Comparison",
    "Demand Summary (APS / Oracle)",
    "Demand Summary (IBP)",
    "Velocity Analysis",
])
def test_each_documented_dpa_section_exists(section):
    """Every section the Demand Planner guide names must be in that view."""
    titles = " | ".join(_expander_titles(_DPA))
    assert section in titles, f"documented but missing from the view: {section}"


def test_the_four_step_table_matches_the_real_step_labels():
    """The primer's step names are quoted from the UI — keep them true."""
    titles = " | ".join(_expander_titles(_DPA))
    for module, s1, s2, s3, s4 in doc._MODULE_STEPS:
        for step in (s1, s2, s3, s4):
            # Step 4 of RO is summarised rather than quoted verbatim.
            if step == "Re-upload / change how RO is read":
                continue
            assert step in titles, f"{module}: no step titled {step!r}"


def test_the_primer_covers_exactly_the_three_upload_modules():
    assert len(doc._MODULE_STEPS) == 3
    assert len(doc._FOUR_STEP_PRIMER) == 4
    names = {m[0] for m in doc._MODULE_STEPS}
    assert names == {"RO Comparison", "Demand Summary (APS / Oracle)",
                     "Demand Summary (IBP)"}


# ── Formulas ────────────────────────────────────────────────────────────────

def test_every_formula_names_a_real_page():
    pages = set(_sidebar_order())
    for f in doc._FORMULAS:
        head = f.page.split("→")[0].strip()
        assert head in pages, f"{f.title}: unknown page {head!r}"


def test_every_formula_has_a_body_and_its_inputs():
    for f in doc._FORMULAS:
        assert f.body.strip() and f.inputs.strip(), f.title


def test_the_bias_formula_points_at_the_step_that_owns_it():
    """It moved into Demand Summary (IBP) Step 3; it used to say Business Health."""
    bias = next(f for f in doc._FORMULAS if "bias" in f.title.lower())
    assert "Demand Summary (IBP)" in bias.page and "Step 3" in bias.page


def test_the_corp_group_formula_states_the_automatic_rule():
    """The retired manual patch is now a documented deterministic fallback."""
    f = next(f for f in doc._FORMULAS if "corporate group" in f.title.lower())
    assert "customer" in f.body.lower()
    assert "fuzzy" in f.notes.lower(), "the fuzzy carve-out must be explained"


# ── The file index ─────────────────────────────────────────────────────────

def test_every_file_is_attached_to_a_real_page():
    """The index is keyed by app page — a typo would orphan a whole group."""
    pages = set(_sidebar_order())
    for f in doc._DATA_FILES:
        assert f.view in pages, f"{f.name}: unknown page {f.view!r}"


def test_the_index_covers_every_page_that_has_files():
    """Not just Demand Planner Analytics — the old version only had that one."""
    covered = {f.view for f in doc._DATA_FILES}
    assert covered >= {
        "Demand Planner Analytics", "Market Barometer", "New Price Quote",
        "Shipment Monitor & HTST Requote", "Bid Assistant",
        "Oracle Pricing Data Download", "Pricing Execution Automation",
    }


def test_every_row_is_complete_and_uses_a_known_owner():
    known = {label for label, _meaning in doc._OWNERS}
    for f in doc._DATA_FILES:
        assert f.name.strip() and f.purpose.strip(), f.name
        assert f.owner in known, f"{f.name}: unknown owner {f.owner!r}"
        assert f.purpose.rstrip().endswith("."), f"{f.name}: purpose not a sentence"


def test_folder_links_are_escaped_deep_links():
    for f in doc._DATA_FILES:
        if not f.folder:
            continue                     # a lakehouse table, not a file
        url = doc._lakehouse_url(f.folder)
        assert url.startswith("https://app.fabric.microsoft.com/")
        assert "selectedPath=Files%2F" in url, f.name
        assert " " not in url, f"unescaped space in {f.name}"


def test_a_table_row_carries_no_folder_link():
    """Delta tables have no folder to open; the row must say so, not link."""
    tables = [f for f in doc._DATA_FILES if not f.folder]
    assert tables, "expected some lakehouse-table rows"
    for f in tables:
        assert "table" in f.name.lower(), f"{f.name}: no folder but not a table"


def test_documented_paths_are_referenced_by_the_code():
    """Each folder or file named here must appear somewhere in data_sources.

    Guards against documenting a path that was renamed or never existed —
    the failure mode where a planner replaces a file nothing reads.
    """
    # Excludes documentation_view itself — it declares these very strings, so
    # including it would make every assertion below trivially true.
    sources = chr(10).join(
        path.read_text(encoding="utf-8")
        for pkg in ("data_sources", "pricebook", "pages", "utils")
        for path in (_ROOT / pkg).glob("*.py")
        if path.name != "documentation_view.py")
    # Folders the app only ever links to — they have no reader of their own.
    link_only = {"Finance", "Monthly_Pricing_Execution", "VBCS"}
    for f in doc._DATA_FILES:
        if not f.folder or f.folder in link_only:
            continue
        assert f.folder in sources, f"{f.name}: folder {f.folder!r} unused"


def test_grouping_keeps_each_page_navigable():
    """A page with many files must be broken into groups, not one long list."""
    for view in doc._DATA_VIEW_ORDER:
        groups = doc._files_for(view)
        count = sum(len(rows) for _g, rows in groups)
        if count > 10:
            assert len(groups) > 1, f"{view}: {count} rows in one flat group"
        # A group is never emitted empty.
        assert all(rows for _g, rows in groups), view


def test_the_file_table_is_one_markdown_block():
    """Row-per-st.markdown renders a stack of one-row tables, not a table."""
    rows = doc._files_for("Bid Assistant")[0][1]
    table = doc._file_table(rows)
    assert table.count("\n") == len(rows) + 1      # header + separator + rows
    assert table.startswith("| File |")


def test_the_replace_instructions_tell_people_to_back_up_and_keep_the_name():
    joined = " ".join(doc._REPLACE_STEPS).lower()
    assert "download" in joined, "no back-up-first step"
    assert "name" in joined and "header" in joined, "no keep-the-schema step"
    assert "app" in doc._REPLACE_WARNING.lower()


# ── The Word export ─────────────────────────────────────────────────────────

@pytest.mark.skipif(not docx_export.AVAILABLE, reason="python-docx not installed")
def test_the_word_export_opens_and_carries_every_section():
    import io
    import docx

    data = _docx_bytes()
    assert data[:2] == b"PK", "a .docx is a zip"
    document = docx.Document(io.BytesIO(data))

    h1 = [p.text for p in document.paragraphs if p.style.name == "Heading 1"]
    assert h1 == [
        "Start here",
        "What each page does, and how to drive it",
        "The three upload modules all work the same way",
        "Every formula, so you can rebuild it in Excel",
        "Where the data lives",
    ]
    h2 = {p.text for p in document.paragraphs if p.style.name == "Heading 2"}
    for g in doc._PAGE_GUIDES:
        assert g.name in h2, f"missing page guide: {g.name}"
    for f in doc._FORMULAS:
        assert f.title in h2, f"missing formula: {f.title}"


@pytest.mark.skipif(not docx_export.AVAILABLE, reason="python-docx not installed")
def test_the_word_export_keeps_the_links_clickable():
    """A manual you cannot click through is much less useful when shared."""
    import io
    import zipfile

    import html

    data = _docx_bytes()
    rels = zipfile.ZipFile(io.BytesIO(data)).read(
        "word/_rels/document.xml.rels").decode()
    # Targets are XML-escaped in the rels part, so a URL carrying a query
    # separator arrives as "&amp;" — unescape before comparing to the source.
    targets = [html.unescape(t) for t in
               re.findall(r'Target="([^"]+)"[^>]*TargetMode="External"', rels)]
    linked = {f.folder for f in doc._DATA_FILES if f.folder}
    assert sum("lakehouses" in t for t in targets) >= len(linked)
    assert any(doc._VELOCITY_REPORT_URL == t for t in targets)
    assert any(doc._FINANCE_PNL_REPORT_URL == t for t in targets)


@pytest.mark.skipif(not docx_export.AVAILABLE, reason="python-docx not installed")
def test_the_word_export_preserves_formula_line_breaks():
    """A formula flattened onto one line is unreadable and untestable."""
    import io
    import docx

    document = docx.Document(io.BytesIO(_docx_bytes()))
    mono = [p for p in document.paragraphs
            if p.runs and p.runs[0].font.name == "Consolas"]
    assert len(mono) == len(doc._FORMULAS)
    wmape = next(p for p in mono if "WMAPE" in p.text)
    assert wmape.text.count("\n") >= 5, "the code block collapsed to one line"


def test_the_page_survives_a_missing_python_docx(monkeypatch):
    """The landing page is the only sign-in surface — it must always render."""
    captions = []
    monkeypatch.setattr(doc.docx_export, "AVAILABLE", False)
    monkeypatch.setattr(doc.st, "caption", lambda t="", **k: captions.append(str(t)))
    monkeypatch.setattr(doc.st, "download_button",
                        lambda *a, **k: pytest.fail("offered a broken download"))
    doc._render_docx_download()
    assert captions and "unavailable" in captions[0]


def test_a_failing_export_degrades_to_a_caption(monkeypatch):
    captions = []
    monkeypatch.setattr(doc.docx_export, "AVAILABLE", True)
    monkeypatch.setattr(doc, "_build_manual_docx",
                        lambda: (_ for _ in ()).throw(RuntimeError("boom")))
    monkeypatch.setattr(doc.st, "caption", lambda t="", **k: captions.append(str(t)))
    monkeypatch.setattr(doc.st, "download_button",
                        lambda *a, **k: pytest.fail("offered a broken download"))
    doc._render_docx_download()
    assert captions and "boom" in captions[0]


def test_docx_builder_raises_clearly_without_the_library(monkeypatch):
    monkeypatch.setattr(docx_export, "AVAILABLE", False)
    with pytest.raises(docx_export.DocxUnavailable, match="python-docx"):
        docx_export.DocBuilder("x")
